"""
python-docx paragraph editor for ScaleXPro CV-Modifier.
Allows viewing and modifying paragraphs in source DOCX files prior to
conversion, redaction, and branding.
"""

from pathlib import Path
from typing import List, Dict, Any, Tuple
import docx
from app.logger import logger


def get_docx_paragraphs(docx_path: str) -> List[Dict[str, Any]]:
    """
    Extracts all text paragraphs from a DOCX file with their 0-based indices.
    """
    if not Path(docx_path).exists():
        return []

    doc = docx.Document(docx_path)
    paragraphs = []
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:  # only return non-empty paragraphs
            paragraphs.append({
                "index": idx,
                "text": text,
                "style": p.style.name if p.style else "Normal",
            })
    return paragraphs


def save_docx_paragraph_edits(
    docx_path: str,
    output_docx_path: str,
    edits: Dict[int, str],
) -> Tuple[bool, str]:
    """
    Applies paragraph edits to a DOCX document and saves the modified version.
    edits is a mapping of {paragraph_index: new_text}.
    """
    try:
        doc = docx.Document(docx_path)
        for idx, new_text in edits.items():
            if 0 <= idx < len(doc.paragraphs):
                p = doc.paragraphs[idx]
                p.text = new_text

        Path(output_docx_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_docx_path)
        return True, "DOCX paragraphs updated successfully."
    except Exception as e:
        logger.error(f"Error saving DOCX paragraph edits: {e}")
        return False, str(e)
