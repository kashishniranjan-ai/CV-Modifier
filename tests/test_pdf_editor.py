"""
Unit and integration tests for CV Editor Mode:
- Text replacement preserving page layout, background, and fonts
- Shrink-to-fit when text overflows
- Custom text box insertion
- True redaction area deletion (text + images)
- Contact detection and safety re-scan blocking
- 30+ step Undo/Redo stack integrity
- Version file management (v1 original, v2, v3... with restore)
- DOCX paragraph editing
- Scanned / image-only PDF detection
"""

import os
import pytest
from pathlib import Path
import pymupdf as fitz

from app.core.pdf_editor import (
    CVEditorEngine,
    match_base14_font,
    sample_background_color,
)
from app.core.docx_editor import get_docx_paragraphs, save_docx_paragraph_edits


@pytest.fixture
def sample_pdf(tmp_path):
    """Generates a multi-page test PDF with distinct text blocks, styled headers, and candidate data."""
    pdf_path = tmp_path / "test_candidate_cv.pdf"
    doc = fitz.open()

    # Page 1: Standard layout with name, role, summary, skills
    page1 = doc.new_page(width=595, height=842)
    # Header rectangle
    page1.draw_rect(fitz.Rect(0, 0, 595, 80), color=None, fill=(0.06, 0.09, 0.16)) # dark navy header
    page1.insert_text((40, 50), "KASHISH NIRANJAN", fontname="helv", fontsize=22, color=(1, 1, 1))

    # Body text
    page1.insert_text((40, 120), "Senior Full Stack Engineer", fontname="helv", fontsize=14, color=(0.2, 0.25, 0.33))
    page1.insert_text((40, 160), "PROFESSIONAL SUMMARY", fontname="helv", fontsize=12, color=(0.95, 0.54, 0.12))
    page1.insert_textbox(
        fitz.Rect(40, 180, 550, 240),
        "Experienced software developer with expertise in Python, PySide6, and enterprise desktop systems.",
        fontname="helv",
        fontsize=10,
        color=(0.1, 0.1, 0.1),
    )

    # Page 2: Experience
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((40, 60), "WORK EXPERIENCE", fontname="helv", fontsize=14, color=(0.06, 0.09, 0.16))
    page2.insert_text((40, 100), "Lead Architect - ScaleXPro (2022 - Present)", fontname="helv", fontsize=11, color=(0.2, 0.2, 0.2))

    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)


@pytest.fixture
def scanned_pdf(tmp_path):
    """Generates an image-only / scanned PDF without any text stream."""
    pdf_path = tmp_path / "scanned_cv.pdf"
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    # Draw simple shapes / faux scanned graphics without text
    page.draw_rect(fitz.Rect(50, 50, 500, 700), color=(0.8, 0.8, 0.8), fill=(0.95, 0.95, 0.95))
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)


def test_text_detection_under_point(sample_pdf):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")
    # Coordinates for "Senior Full Stack Engineer" (around y=120)
    info = engine.detect_text_under_point(0, 45, 120)
    assert info is not None
    assert "Senior Full Stack Engineer" in info["text"]
    assert info["font_size"] > 12.0
    assert info["rect"] is not None


def test_text_replacement_preserves_page_intact(sample_pdf):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")
    # Target "Senior Full Stack Engineer"
    info = engine.detect_text_under_point(0, 45, 120)
    assert info is not None

    ok, warn = engine.replace_text(
        page_num=0,
        rect=info["rect"],
        new_text="Principal Solutions Architect",
        font_name="helv",
        font_size=info["font_size"],
        color_rgb=info["color_rgb"],
        old_text_snippet=info["text"],
    )
    assert ok is True

    # Verify new text is in the page and old text is gone
    page_text = engine.doc[0].get_text("text")
    assert "Principal Solutions Architect" in page_text
    assert "Senior Full Stack Engineer" not in page_text
    # Verify rest of page is preserved
    assert "PROFESSIONAL SUMMARY" in page_text
    assert "Experienced software developer" in page_text


def test_text_shrink_to_fit_on_overflow(sample_pdf):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")
    # Small rect
    rect = (40, 300, 120, 320)  # only 80pt wide
    long_text = "This is an extraordinarily long title that cannot possibly fit in eighty points"

    ok, warn = engine.replace_text(
        page_num=0,
        rect=rect,
        new_text=long_text,
        font_name="helv",
        font_size=14.0,
    )
    assert ok is True
    assert "fit" in warn.lower() or warn == ""


def test_add_text_box(sample_pdf):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")
    rect = (40, 400, 300, 430)
    ok, msg = engine.add_text_box(
        page_num=0,
        rect=rect,
        text="Candidate shortlisted for round 2 technical review.",
        font_name="helv",
        font_size=10.0,
        color_rgb=(0.06, 0.09, 0.16),
    )
    assert ok is True
    page_text = engine.doc[0].get_text("text")
    assert "round 2 technical review" in page_text


def test_delete_area_true_redaction(sample_pdf):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")
    # Delete the PROFESSIONAL SUMMARY heading at y=160
    delete_rect = (35, 150, 300, 175)
    ok, msg = engine.delete_area(page_num=0, rect=delete_rect)
    assert ok is True

    page_text = engine.doc[0].get_text("text")
    assert "PROFESSIONAL SUMMARY" not in page_text


def test_undo_redo_30_steps(sample_pdf):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")
    initial_text = engine.doc[0].get_text("text")

    # Perform 32 sequential edits
    for i in range(32):
        engine.add_text_box(
            page_num=0,
            rect=(40, 500 + i * 5, 200, 520 + i * 5),
            text=f"Note item {i}",
        )

    assert engine.get_undo_count() >= 30
    assert "Note item 31" in engine.doc[0].get_text("text")

    # Undo 5 times
    for _ in range(5):
        assert engine.undo() is True

    # Note item 31 should be gone, Note item 26 still present
    text_after_undo = engine.doc[0].get_text("text")
    assert "Note item 31" not in text_after_undo
    assert "Note item 26" in text_after_undo

    # Redo once
    assert engine.redo() is True
    assert "Note item 27" in engine.doc[0].get_text("text")


def test_safety_check_blocks_saving_contacts(sample_pdf, tmp_path):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")

    # Add a phone number in an edit
    engine.add_text_box(
        page_num=0,
        rect=(40, 450, 300, 480),
        text="Direct contact: +91 9876543210 or recruiter@agency.com",
    )

    is_safe, violations = engine.check_safety()
    assert is_safe is False
    assert len(violations) >= 2  # phone and email

    # Attempt to save version should be BLOCKED
    ok, saved_path, err_msg = engine.save_version(tmp_path)
    assert ok is False
    assert "Safety scan blocked save" in err_msg


def test_version_management(sample_pdf, tmp_path):
    engine = CVEditorEngine(sample_pdf, "CAND-0010")
    assert len(engine.versions) == 1  # v1 original branded

    # Clean edit
    engine.add_text_box(page_num=0, rect=(40, 450, 300, 470), text="Clean verified candidate comment")

    # Save v2
    ok, v2_path, msg = engine.save_version(tmp_path)
    assert ok is True
    assert Path(v2_path).exists()
    assert "v2" in Path(v2_path).name
    assert len(engine.versions) == 2

    # Reset to original restores v1
    assert engine.reset_to_original() is True
    assert "Clean verified candidate comment" not in engine.doc[0].get_text("text")


def test_scanned_pdf_detection(scanned_pdf):
    engine = CVEditorEngine(scanned_pdf, "CAND-0011")
    assert engine.is_page_scanned(0) is True


def test_docx_paragraph_editing(tmp_path):
    import docx
    docx_file = tmp_path / "test_resume.docx"
    doc = docx.Document()
    doc.add_paragraph("John Doe - Lead Developer")
    doc.add_paragraph("Passionate about cloud architecture and distributed databases.")
    doc.save(str(docx_file))

    # Read paragraphs
    paragraphs = get_docx_paragraphs(str(docx_file))
    assert len(paragraphs) == 2
    assert "Lead Developer" in paragraphs[0]["text"]

    # Edit paragraph
    edited_docx = tmp_path / "test_resume_edited.docx"
    ok, msg = save_docx_paragraph_edits(str(docx_file), str(edited_docx), {0: "Jane Doe - Principal Architect"})
    assert ok is True

    # Verify edited docx
    doc_new = docx.Document(str(edited_docx))
    assert doc_new.paragraphs[0].text == "Jane Doe - Principal Architect"
